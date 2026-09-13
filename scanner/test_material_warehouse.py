from decimal import Decimal
from io import BytesIO

import openpyxl

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile

from scanner.material_models import (
    AuxiliaryMaterialLot,
    AuxiliaryMaterialTransaction,
    MaterialGrade,
    MaterialRequirement,
    MaterialStockLot,
    MaterialTransaction,
)
from scanner.material_services import create_material_request, issue_request_lines
from scanner.material_stock_import import (
    create_stock_lots_from_preview,
    group_auxiliary_lots,
    group_stock_lots,
    parse_stock_workbook,
    reconcile_stock_from_preview,
)
from scanner.models import Employee, Order


class MaterialCalculationTests(TestCase):
    def setUp(self):
        self.steel = MaterialGrade.objects.create(name="Тестовая сталь", category="steel", density_kg_m3=7850)

    def test_sheet_area_and_mass(self):
        item = MaterialRequirement(
            grade=self.steel,
            profile_type="sheet",
            thickness_mm=8,
            width_mm=1000,
            piece_length_mm=2000,
        )
        self.assertEqual(item.calculate_area_m2(2), Decimal("4.000"))
        self.assertEqual(item.calculate_mass(2, 0), Decimal("251.200"))

    def test_round_pipe_mass_by_total_remaining_length(self):
        item = MaterialRequirement(
            grade=self.steel,
            profile_type="round_pipe",
            outer_diameter_mm=57,
            wall_thickness_mm=3.5,
        )
        self.assertEqual(item.calculate_mass(0, 12000), Decimal("55.414"))

    def test_square_and_hexagon_mass_per_meter(self):
        square = MaterialRequirement(
            grade=self.steel, profile_type="square_bar", width_mm=20,
        )
        hexagon = MaterialRequirement(
            grade=self.steel, profile_type="hex_bar", width_mm=24,
        )
        self.assertEqual(square.calculate_mass(0, 1000), Decimal("3.140"))
        self.assertEqual(hexagon.calculate_mass(0, 1000), Decimal("3.916"))


class MaterialIssueTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="material-test", password="x")
        Employee.objects.bulk_create([Employee(last_name="Иванов", first_name="Иван", is_active=True)])
        self.employee = Employee.objects.get(last_name="Иванов")
        self.order = Order.objects.create(order_number="TEST-1", full_name="Тестовый проект")
        self.steel = MaterialGrade.objects.create(name="Сталь для выдачи", category="steel", density_kg_m3=7850)
        self.requirement = MaterialRequirement.objects.create(
            order=self.order,
            item_name="Труба на раму",
            grade=self.steel,
            profile_type="round_pipe",
            outer_diameter_mm=57,
            wall_thickness_mm=3.5,
            piece_length_mm=3000,
            quantity_required=2,
            total_length_required_mm=6000,
        )
        self.lot = MaterialStockLot(
            order=self.order,
            name="Труба 57×3,5",
            grade=self.steel,
            profile_type="round_pipe",
            outer_diameter_mm=57,
            wall_thickness_mm=3.5,
            piece_length_mm=6000,
            quantity_initial=2,
            length_initial_mm=12000,
            created_by=self.user,
        )
        self.lot.initialize_balances()
        self.lot.save()

    def test_partial_issue_updates_length_mass_and_document(self):
        document = create_material_request(self.order, [self.requirement], self.user, "Цех 1")
        line = document.lines.get()
        token, rows = issue_request_lines(
            document,
            {line.id: {"length_mm": "3000", "quantity": "0"}},
            self.employee,
            "На раму",
            self.user,
        )
        self.assertTrue(token)
        self.assertEqual(len(rows), 1)
        self.lot.refresh_from_db()
        document.refresh_from_db()
        line.refresh_from_db()
        self.assertEqual(self.lot.length_remaining_mm, Decimal("9000"))
        self.assertEqual(line.length_issued_mm, Decimal("3000"))
        self.assertEqual(document.status, "partial")
        self.assertEqual(MaterialTransaction.objects.filter(transaction_type="out").count(), 1)

    def test_overissue_rolls_back_stock(self):
        document = create_material_request(self.order, [self.requirement], self.user)
        line = document.lines.get()
        with self.assertRaises(ValidationError):
            issue_request_lines(
                document,
                {line.id: {"length_mm": "7000", "quantity": "0"}},
                self.employee,
                "",
                self.user,
            )
        self.lot.refresh_from_db()
        self.assertEqual(self.lot.length_remaining_mm, Decimal("12000"))
        self.assertFalse(MaterialTransaction.objects.filter(transaction_type="out").exists())


class MaterialStockImportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="stock-import", password="x")
        self.order = Order.objects.create(order_number="STOCK-PROJECT", full_name="Проект склада")
        self.steel = MaterialGrade.objects.create(
            name="Сталь импорт", category="steel", density_kg_m3=7850
        )

    def workbook(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Труба"
        sheet.append([
            "Наименование", "Вид трубы", "Марка материала", "Сортамент",
            "Наружный диаметр, мм", "Толщина стенки, мм", "Длина куска, мм",
            "Количество", "Партия / плавка", "Место хранения",
        ])
        sheet.append(["Труба 57×3,5", "Труба круглая", self.steel.name, "57×3,5", 57, 3.5, 6000, 2, "П-1", "Т-1"])
        sheet.append(["Труба 57×3,5", "Труба круглая", self.steel.name, "57×3,5", 57, 3.5, 3500, 1, "П-1", "Т-1"])
        return workbook

    def test_preview_and_confirm_split_linear_stock_into_pieces(self):
        rows, errors = parse_stock_workbook(self.workbook())
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)
        self.assertEqual(sum(row["lots_to_create"] for row in rows), 3)

        lots = create_stock_lots_from_preview(rows, self.user)
        self.assertEqual(len(lots), 3)
        self.assertEqual(MaterialTransaction.objects.filter(transaction_type="in").count(), 3)
        self.assertEqual(
            sum((lot.length_remaining_mm for lot in lots), Decimal("0")),
            Decimal("15500"),
        )

        groups = group_stock_lots(
            MaterialStockLot.objects.select_related("grade", "order").all()
        )
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["quantity"], Decimal("3"))
        self.assertEqual(groups[0]["length_m"], Decimal("15.500"))
        self.assertEqual(len(groups[0]["details"]), 2)

    def test_unknown_grade_blocks_whole_preview(self):
        workbook = self.workbook()
        workbook["Труба"]["C2"] = "Неизвестная марка"
        rows, errors = parse_stock_workbook(workbook)
        self.assertEqual(len(rows), 1)
        self.assertTrue(any("отсутствует в справочнике" in error for error in errors))

    def test_rolled_profile_names_include_hexagon_beam_and_bulb_flat(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Прокат"
        sheet.append([
            "Наименование", "Вид профиля", "Марка материала", "Сортамент",
            "Ширина, мм", "Длина куска, мм", "Количество", "Масса 1 м, кг",
        ])
        sheet.append(["Шестигранник 24", "Шестигранник", self.steel.name, "S24", 24, 3000, 2, ""])
        sheet.append(["Швеллер 12П", "Швеллер", self.steel.name, "12П", "", 6000, 1, 10.4])
        sheet.append(["Двутавр 20Б1", "Двутавр", self.steel.name, "20Б1", "", 6000, 1, 21.3])
        sheet.append(["Полособульб", "Полособульб", self.steel.name, "100×8", "", 6000, 1, 8.5])
        rows, errors = parse_stock_workbook(workbook)
        self.assertEqual(errors, [])
        self.assertEqual(
            [row["profile_type"] for row in rows],
            ["hex_bar", "channel", "beam", "bulb_flat"],
        )

    def test_stock_pages_and_template_open(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get("/materials/stock/").status_code, 200)
        self.assertEqual(self.client.get("/materials/stock/import/").status_code, 200)
        response = self.client.get("/materials/stock/import/template/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

    def test_web_preview_then_confirmation(self):
        output = BytesIO()
        self.workbook().save(output)
        upload = SimpleUploadedFile(
            "stock.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.client.force_login(self.user)
        response = self.client.post(
            "/materials/stock/import/",
            {"action": "preview", "file": upload},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Подтвердить загрузку")
        session_keys = [
            key for key in self.client.session.keys()
            if key.startswith("material_stock_import:")
        ]
        self.assertEqual(len(session_keys), 1)
        token = session_keys[0].split(":", 1)[1]
        response = self.client.post(
            "/materials/stock/import/",
            {"action": "confirm", "token": token},
        )
        self.assertRedirects(response, "/materials/stock/")
        self.assertEqual(MaterialStockLot.objects.count(), 3)

    def test_web_selection_assigns_one_project_and_basis_to_whole_file(self):
        output = BytesIO()
        self.workbook().save(output)
        upload = SimpleUploadedFile(
            "stock.xlsx", output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.client.force_login(self.user)
        response = self.client.post(
            "/materials/stock/import/",
            {
                "action": "preview", "file": upload,
                "order_id": str(self.order.id), "basis": "Накладная № 125",
            },
        )
        self.assertEqual(response.status_code, 200)
        session_key = next(
            key for key in self.client.session.keys()
            if key.startswith("material_stock_import:")
        )
        rows = self.client.session[session_key]
        self.assertTrue(all(row["order_id"] == self.order.id for row in rows))
        self.assertTrue(all(row["basis"] == "Накладная № 125" for row in rows))

    def test_inventory_upload_uses_general_stock_and_inventory_basis(self):
        output = BytesIO()
        self.workbook().save(output)
        upload = SimpleUploadedFile(
            "inventory.xlsx", output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.client.force_login(self.user)
        response = self.client.post(
            "/materials/stock/import/",
            {"action": "preview", "file": upload, "receipt_type": "inventory"},
        )
        self.assertEqual(response.status_code, 200)
        session_key = next(
            key for key in self.client.session.keys()
            if key.startswith("material_stock_import:")
        )
        rows = self.client.session[session_key]
        self.assertTrue(all(row["order_id"] is None for row in rows))
        self.assertTrue(all(row["basis"] == "Остатки после инвентаризации" for row in rows))

    def test_repeated_inventory_replaces_active_balance_without_doubling(self):
        rows, errors = parse_stock_workbook(self.workbook())
        self.assertEqual(errors, [])
        for row in rows:
            row["basis"] = "Инвентаризация"

        created, adjusted = reconcile_stock_from_preview(rows, self.user)
        self.assertEqual(len(created), 3)
        self.assertEqual(adjusted, 0)
        self.assertEqual(
            sum(
                MaterialStockLot.objects.values_list("length_remaining_mm", flat=True),
                Decimal("0"),
            ),
            Decimal("15500"),
        )

        created, adjusted = reconcile_stock_from_preview(rows, self.user)
        self.assertEqual(len(created), 3)
        self.assertEqual(adjusted, 3)
        self.assertEqual(MaterialStockLot.objects.filter(length_remaining_mm__gt=0).count(), 3)
        self.assertEqual(MaterialStockLot.objects.count(), 6)
        self.assertEqual(
            sum(
                MaterialStockLot.objects.values_list("length_remaining_mm", flat=True),
                Decimal("0"),
            ),
            Decimal("15500"),
        )

    def test_actual_stock_export_round_trips_without_doubling(self):
        rows, errors = parse_stock_workbook(self.workbook())
        self.assertEqual(errors, [])
        create_stock_lots_from_preview(rows, self.user)
        self.client.force_login(self.user)
        response = self.client.get("/materials/stock/export/")
        self.assertEqual(response.status_code, 200)
        exported = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(
            exported.sheetnames,
            ["Инструкция", "Лист", "Труба", "Прокат", "Прочие материалы"],
        )
        exported_rows, export_errors = parse_stock_workbook(exported)
        self.assertEqual(export_errors, [])
        self.assertEqual(len(exported_rows), 2)
        for row in exported_rows:
            row["basis"] = "Повторная инвентаризация"
            row["inventory_sections"] = {"metal": True, "auxiliary": True}
        reconcile_stock_from_preview(exported_rows, self.user)
        self.assertEqual(MaterialStockLot.objects.filter(length_remaining_mm__gt=0).count(), 3)
        self.assertEqual(
            sum(
                MaterialStockLot.objects.values_list("length_remaining_mm", flat=True),
                Decimal("0"),
            ),
            Decimal("15500"),
        )

class AuxiliaryMaterialTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="auxiliary-stock", password="x")
        self.employee = Employee.objects.create(last_name="Петров", first_name="Пётр", is_active=True)

    def workbook(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Прочие материалы"
        sheet.append([
            "Категория", "Наименование", "Марка / производитель", "Характеристика",
            "Единица учета", "Количество", "Тара / упаковка", "Плотность, кг/л",
            "Толщина, мм", "Ширина, мм", "Длина, мм", "Ячейка X, мм", "Ячейка Y, мм",
            "Диаметр проволоки, мм", "Партия", "Срок годности", "Место хранения",
            "Проект", "ЛВЖ / опасный", "Минимальный остаток", "Основание прихода",
        ])
        sheet.append([
            "Краска / покрытие", "Эмаль ПФ-115", "Лакра", "RAL 5005", "л", 20,
            "4 банки по 5 л", 1.2, "", "", "", "", "", "", "П-1", "31.12.2027",
            "Шкаф ЛВЖ", "", "Да", 5, "Инвентаризация",
        ])
        sheet.append([
            "Сетка", "Сетка сварная", "", "Карта сетки", "м²", 12, "", "", "",
            1000, 2000, 50, 50, 3, "С-1", "", "Стеллаж С-1", "", "Нет", 2,
            "Инвентаризация",
        ])
        return workbook

    def test_preview_confirm_mass_mesh_and_grouping(self):
        rows, errors = parse_stock_workbook(self.workbook())
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["mass_kg"], "24.000")
        self.assertEqual(rows[1]["area_m2"], "12.000")

        lots = create_stock_lots_from_preview(rows, self.user)
        self.assertEqual(len(lots), 2)
        self.assertEqual(AuxiliaryMaterialTransaction.objects.filter(transaction_type="in").count(), 2)
        paint = AuxiliaryMaterialLot.objects.get(category="paint")
        mesh = AuxiliaryMaterialLot.objects.get(category="mesh")
        self.assertEqual(paint.estimated_mass_kg, Decimal("24.000"))
        self.assertIn("ячейка 50.000 × 50.000 мм", mesh.dimensions_display)
        groups = group_auxiliary_lots(AuxiliaryMaterialLot.objects.all())
        self.assertEqual(len(groups), 2)

    def test_manual_issue_subtracts_and_overissue_is_rejected(self):
        lot = AuxiliaryMaterialLot.objects.create(
            category="lubricant", name="Смазка", unit="kg",
            quantity_initial=10, quantity_remaining=10, created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.post(
            f"/materials/stock/auxiliary/{lot.id}/issue/",
            {"quantity": "3.5", "recipient_id": self.employee.id, "basis": "ТО оборудования"},
        )
        self.assertRedirects(response, "/materials/stock/")
        lot.refresh_from_db()
        self.assertEqual(lot.quantity_remaining, Decimal("6.500"))
        movement = AuxiliaryMaterialTransaction.objects.get(transaction_type="out")
        self.assertEqual(movement.recipient_name, "Петров Пётр")

        response = self.client.post(
            f"/materials/stock/auxiliary/{lot.id}/issue/",
            {"quantity": "7", "recipient_id": self.employee.id},
        )
        self.assertRedirects(response, "/materials/stock/")
        lot.refresh_from_db()
        self.assertEqual(lot.quantity_remaining, Decimal("6.500"))
        self.assertEqual(AuxiliaryMaterialTransaction.objects.filter(transaction_type="out").count(), 1)

    def test_repeated_auxiliary_inventory_replaces_balance(self):
        rows, errors = parse_stock_workbook(self.workbook())
        self.assertEqual(errors, [])
        for row in rows:
            row["basis"] = "Инвентаризация"
        reconcile_stock_from_preview(rows, self.user)
        created, adjusted = reconcile_stock_from_preview(rows, self.user)
        self.assertEqual(len(created), 2)
        self.assertEqual(adjusted, 2)
        self.assertEqual(AuxiliaryMaterialLot.objects.filter(quantity_remaining__gt=0).count(), 2)
        self.assertEqual(AuxiliaryMaterialLot.objects.count(), 4)

    def test_auxiliary_actual_stock_export_can_be_imported(self):
        rows, errors = parse_stock_workbook(self.workbook())
        self.assertEqual(errors, [])
        create_stock_lots_from_preview(rows, self.user)
        self.client.force_login(self.user)
        response = self.client.get("/materials/stock/export/")
        self.assertEqual(response.status_code, 200)
        exported = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
        exported_rows, export_errors = parse_stock_workbook(exported)
        self.assertEqual(export_errors, [])
        self.assertEqual(len(exported_rows), 2)
        self.assertTrue(all(row["record_type"] == "auxiliary" for row in exported_rows))
