from io import BytesIO

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from scanner.models import Order
from scanner.purchase_import import enrich_purchase_preview, parse_purchase_workbook
from scanner.purchase_allocation import allocation_state
from scanner.purchase_models import (
    PurchaseItem, PurchasePreparation, PurchaseRequest, PurchaseRequestLine,
    PurchaseTransaction,
)
from scanner.purchase_normalization import normalize_purchase_name


class PurchaseNameNormalizationTests(TestCase):
    def test_harmless_punctuation_and_spacing_are_equal(self):
        base = "Болт М10-6gx25.88.0118 ГОСТ 7798-70"
        variants = [
            "  Болт  М10–6g×25.88.0118  ГОСТ 7798-70. ",
            "болт м10-6Gх25.88.0118 гост 7798-70",
            "Болт М10 - 6gx25.88.0118 ГОСТ. 7798-70",
        ]
        for value in variants:
            self.assertEqual(normalize_purchase_name(value), normalize_purchase_name(base))

    def test_decimal_point_is_not_discarded(self):
        self.assertNotEqual(normalize_purchase_name("Шаг 1.25"), normalize_purchase_name("Шаг 125"))

    def test_dot_variant_uses_one_shared_stock_balance(self):
        order = Order.objects.create(order_number="NORMALIZED-STOCK")
        first = PurchaseItem.objects.create(order=order, item_name="Болт ГОСТ 7798-70", assembly_name="Узел 1", quantity_required=5, quantity_purchased=5)
        second = PurchaseItem.objects.create(order=order, item_name="Болт ГОСТ 7798-70.", assembly_name="Узел 2", quantity_required=5, quantity_purchased=5)
        PurchaseTransaction.objects.create(purchase_item=first, transaction_type="out", quantity=2)
        self.assertEqual(allocation_state(second).purchased_for_order, 5)
        self.assertEqual(allocation_state(second).stock_available, 3)


class PurchaseImportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="purchase-import", password="x")
        self.order = Order.objects.create(order_number="IMPORT-1", full_name="Импорт крепежа")

    def workbook(self, required=2):
        workbook = openpyxl.Workbook()
        instruction = workbook.active
        instruction.title = "Инструкция"
        instruction.append(["Тест"])
        spec = workbook.create_sheet("Спецификация")
        spec.append(["Наименование", "Обозначение / артикул", "Требуемое количество", "Узел / подсборка"])
        spec.append(["Болт М10-6gx25.88.0118 ГОСТ 7798-70", "", required, "Полумуфта"])
        spec.append(["Болт М10-6gx25.88.0118 ГОСТ 7798-70.", "", 1, "Полумуфта"])
        spec.append(["Болт М10-6gx25.88.0118 ГОСТ 7798-70", "", 4, "Корпус"])
        purchase = workbook.create_sheet("Закупка")
        purchase.append(["Наименование", "Закуплено"])
        purchase.append(["Болт М10-6gx25.88.0118 ГОСТ 7798-70.", 20])
        return workbook

    def upload(self, workbook, action="preview"):
        output = BytesIO()
        workbook.save(output)
        return self.client.post("/purchases/import/", {
            "action": action,
            "order_id": self.order.id,
            "file": SimpleUploadedFile("purchase.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        })

    def test_parser_combines_dot_variant_inside_same_assembly(self):
        rows, errors, warnings = parse_purchase_workbook(self.workbook(), self.order)
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)
        polumufta = next(row for row in rows if row["assembly_name"] == "Полумуфта")
        self.assertEqual(polumufta["quantity_required"], 3)
        self.assertEqual(polumufta["quantity_purchased"], 20)
        self.assertTrue(any("Одинаковое наименование" in warning for warning in warnings))

    def test_parser_accepts_decimal_quantities_with_dot_and_comma(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Спецификация"
        sheet.append(["Наименование", "Требуемое количество", "Узел / подсборка"])
        sheet.append(["Рукав 2SN DN10", "0.54", "Узел 1"])
        sheet.append(["Рукав 4SP DN12", "0,71", "Узел 2"])
        rows, errors, _ = parse_purchase_workbook(workbook, self.order)
        self.assertEqual(errors, [])
        self.assertEqual(str(rows[0]["quantity_required"]), "0.54")
        self.assertEqual(str(rows[1]["quantity_required"]), "0.71")

    def test_final_preview_marks_changes_and_calculates_deficit(self):
        PurchaseItem.objects.create(
            order=self.order,
            item_name="Болт М10-6gx25.88.0118 ГОСТ 7798-70",
            assembly_name="Полумуфта",
            quantity_required=3,
            quantity_purchased=5,
        )
        workbook = self.workbook()
        workbook["Закупка"]["B2"] = 5
        rows, errors, _ = parse_purchase_workbook(workbook, self.order)
        preview = enrich_purchase_preview(rows, order=self.order)
        self.assertEqual(errors, [])
        polumufta = next(row for row in preview if row["assembly_name"] == "Полумуфта")
        corpus = next(row for row in preview if row["assembly_name"] == "Корпус")
        self.assertEqual(polumufta["action_code"], "same")
        self.assertEqual(corpus["action_code"], "add")
        self.assertEqual(polumufta["required_after_total"], 7)
        self.assertEqual(polumufta["available_after"], 5)
        self.assertEqual(polumufta["deficit_after"], 2)

    def test_import_page_shows_final_state_columns_and_deficit(self):
        self.client.force_login(self.user)
        workbook = self.workbook()
        workbook["Закупка"]["B2"] = 5
        response = self.upload(workbook)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Итоговый предпросмотр")
        self.assertContains(response, "С корректировкой")
        self.assertContains(response, "Доступно")
        self.assertContains(response, "Дефицит")
        self.assertContains(response, "2")
        self.assertEqual(PurchaseItem.objects.count(), 0)

    def test_preparation_binding_has_preview_before_confirmation(self):
        self.client.force_login(self.user)
        preparation = PurchasePreparation.objects.create(name="Черновик", created_by=self.user)
        PurchaseItem.objects.create(
            preparation=preparation,
            item_name="Болт М12 ГОСТ 7798-70",
            assembly_name="Рама",
            quantity_required=10,
            quantity_purchased=6,
        )
        response = self.client.post(
            f"/purchases/preparation/{preparation.id}/",
            {"action": "preview_bind", "order_id": self.order.id},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Предпросмотр привязки")
        self.assertContains(response, "Подтвердить привязку")
        preparation.refresh_from_db()
        self.assertIsNone(preparation.assigned_order_id)

    def test_statement_management_is_administrator_only(self):
        ordinary = User.objects.create_user(username="ordinary", password="x")
        self.client.force_login(ordinary)
        response = self.client.get("/purchases/manage/")
        self.assertEqual(response.status_code, 403)

    def test_administrator_can_delete_unbound_preparation(self):
        self.client.force_login(self.user)
        preparation = PurchasePreparation.objects.create(name="Удалить", created_by=self.user)
        PurchaseItem.objects.create(
            preparation=preparation, item_name="Гайка М10", quantity_required=4, quantity_purchased=4,
        )
        response = self.client.post(
            f"/purchases/manage/preparation/{preparation.id}/delete/", {"confirm": "yes"},
        )
        self.assertRedirects(response, "/purchases/manage/")
        self.assertFalse(PurchasePreparation.objects.filter(pk=preparation.id).exists())
        self.assertFalse(PurchaseItem.objects.filter(item_name="Гайка М10").exists())

    def test_order_statement_cleanup_removes_unissued_requests_but_keeps_order(self):
        self.client.force_login(self.user)
        item = PurchaseItem.objects.create(
            order=self.order, item_name="Шайба 10", quantity_required=8, quantity_purchased=8,
        )
        document = PurchaseRequest.objects.create(number="TEST-CLEAN-1", order=self.order, requested_by=self.user)
        PurchaseRequestLine.objects.create(request=document, purchase_item=item, quantity_requested=8)
        response = self.client.post(
            f"/purchases/manage/order/{self.order.id}/delete/", {"confirm": "yes"},
        )
        self.assertRedirects(response, "/purchases/manage/")
        self.assertTrue(Order.objects.filter(pk=self.order.id).exists())
        self.assertFalse(PurchaseItem.objects.filter(pk=item.id).exists())
        self.assertFalse(PurchaseRequest.objects.filter(pk=document.id).exists())

    def test_order_statement_with_movements_cannot_be_deleted(self):
        self.client.force_login(self.user)
        item = PurchaseItem.objects.create(
            order=self.order, item_name="Подшипник", quantity_required=2, quantity_purchased=2,
        )
        PurchaseTransaction.objects.create(purchase_item=item, transaction_type="out", quantity=1)
        response = self.client.post(
            f"/purchases/manage/order/{self.order.id}/delete/", {"confirm": "yes"},
        )
        self.assertRedirects(response, "/purchases/manage/")
        self.assertTrue(PurchaseItem.objects.filter(pk=item.id).exists())

    def test_preview_does_not_write_and_confirmation_is_idempotent(self):
        self.client.force_login(self.user)
        response = self.upload(self.workbook())
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Подтвердить загрузку")
        self.assertEqual(PurchaseItem.objects.count(), 0)
        session_key = next(key for key in self.client.session.keys() if key.startswith("purchase_import:"))
        token = session_key.split(":", 1)[1]
        response = self.client.post("/purchases/import/", {"action": "confirm", "token": token})
        self.assertRedirects(response, f"/purchases/{self.order.id}/")
        self.assertEqual(PurchaseItem.objects.count(), 2)
        self.assertEqual(PurchaseItem.objects.get(assembly_name="Полумуфта").quantity_required, 3)

        response = self.upload(self.workbook(required=5))
        token = next(key for key in self.client.session.keys() if key.startswith("purchase_import:")).split(":", 1)[1]
        self.client.post("/purchases/import/", {"action": "confirm", "token": token})
        self.assertEqual(PurchaseItem.objects.count(), 2)
        self.assertEqual(PurchaseItem.objects.get(assembly_name="Полумуфта").quantity_required, 6)

    def test_template_contains_instruction_examples_and_two_input_sheets(self):
        self.client.force_login(self.user)
        response = self.client.get("/purchases/import/template/")
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Инструкция", "Спецификация", "Закупка"])
        self.assertGreater(workbook["Спецификация"].max_row, 1)
        self.assertGreater(workbook["Закупка"].max_row, 1)
