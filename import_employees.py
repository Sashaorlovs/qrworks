import openpyxl, re, os, django, sys
sys.path.insert(0, '/opt/qr_simple')
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'qr_simple.settings')
django.setup()
from scanner.models import Employee

wb = openpyxl.load_workbook('/opt/qr_simple/templates/Список сотрудников.xlsx')
ws = wb.active

pattern = re.compile(r'^([А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.[ А-ЯЁ]?\.?)\s*\((.+)\)$')
cnt = 0

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2, values_only=True):
    cell = row[1]
    if not cell or not isinstance(cell, str):
        continue
    cleaned = cell.replace('\n', ' ').strip()
    m = pattern.match(cleaned)
    if not m:
        continue
    fio_parts = m.group(1).split()
    last_name = fio_parts[0]
    first_name = fio_parts[1] if len(fio_parts) > 1 else ''
    middle_name = fio_parts[2] if len(fio_parts) > 2 else ''
    position = m.group(2)

    emp, created = Employee.objects.get_or_create(
        last_name=last_name,
        first_name=first_name,
        defaults={'middle_name': middle_name, 'position': position, 'is_active': True}
    )
    if not created and emp.position != position:
        emp.position = position
        emp.save(update_fields=['position'])
    if created:
        cnt += 1

print(f'Добавлено новых сотрудников: {cnt}')
print(f'Всего сотрудников в базе: {Employee.objects.count()}')
